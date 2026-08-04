"""src/rd2/extractors/excel.py 테스트.

openpyxl은 리더이자 라이터라 HWP와 달리 합성 fixture를 만들 수 있다 — 실제
수집 문서를 커밋하지 않고 tmp_path에 워크북을 생성해 검증한다.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from rd2.extractors import excel as excel_extractor
from rd2.extractors.excel import ExtractedExcelDocument, extract_excel, render_cell

_OLE_HEADER = bytes.fromhex("d0cf11e0a1b11ae1")


def _write_workbook(path: Path, sheets: list[tuple[str, list[list]]]) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        worksheet = workbook.create_sheet(title=name)
        for row in rows:
            worksheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


class TestNeedsQuarantine:
    def test_normal_document_does_not_need_quarantine(self):
        document = ExtractedExcelDocument(source_path=Path("x.xlsx"))
        assert document.needs_quarantine is False

    def test_encrypted_document_needs_quarantine(self):
        document = ExtractedExcelDocument(source_path=Path("x.xlsx"), is_encrypted=True)
        assert document.needs_quarantine is True

    def test_invalid_document_needs_quarantine(self):
        document = ExtractedExcelDocument(source_path=Path("x.xlsx"), is_valid=False)
        assert document.needs_quarantine is True


class TestRenderCell:
    def test_integral_floats_lose_the_spreadsheet_decimal_point(self):
        assert render_cell(2024.0) == "2024"
        assert render_cell(1.5) == "1.5"

    def test_dates_and_times_are_rendered_as_iso(self):
        assert render_cell(dt.datetime(2024, 6, 6)) == "2024-06-06"
        assert render_cell(dt.datetime(2024, 6, 6, 14, 30, 5)) == "2024-06-06 14:30:05"
        assert render_cell(dt.date(2024, 6, 6)) == "2024-06-06"
        assert render_cell(dt.time(9, 0)) == "09:00:00"

    def test_booleans_and_empty_cells(self):
        assert render_cell(True) == "TRUE"
        assert render_cell(False) == "FALSE"
        assert render_cell(None) == ""

    def test_cell_newlines_are_folded_so_a_row_stays_one_logical_line(self):
        assert render_cell("첫 줄\n둘째 줄\r\n셋째 줄") == "첫 줄 둘째 줄 셋째 줄"


class TestExtractExcel:
    def test_every_worksheet_is_extracted_in_workbook_order(self, tmp_path: Path):
        path = _write_workbook(
            tmp_path / "book.xlsx",
            [
                ("2024년 1차", [["이사", "회차"], ["홍길동", 3]]),
                ("2024년 2차", [["이사"], ["김철수"]]),
            ],
        )

        document = extract_excel(path)

        assert document.is_valid is True
        assert document.needs_quarantine is False
        assert document.truncated is False
        assert [sheet.name for sheet in document.sheets] == ["2024년 1차", "2024년 2차"]
        assert document.sheets[0].rows == [["이사", "회차"], ["홍길동", "3"]]
        assert document.sheets[1].rows == [["이사"], ["김철수"]]

    def test_empty_rows_are_dropped_and_trailing_empty_cells_trimmed(self, tmp_path: Path):
        path = _write_workbook(
            tmp_path / "sparse.xlsx",
            [("Sheet1", [["가", None, None], [None, None, None], [None, "나", None]])],
        )

        document = extract_excel(path)

        assert document.sheets[0].rows == [["가"], ["", "나"]]

    def test_non_zip_input_is_invalid_without_raising(self, tmp_path: Path):
        path = tmp_path / "stub.xlsx"
        path.write_text("<html>파일이 존재하지 않습니다</html>", encoding="utf-8")

        document = extract_excel(path)

        assert document.is_valid is False
        assert document.error == "unknown_excel_format"
        assert document.sheets == []

    def test_ole_container_named_xlsx_is_reported_as_encrypted(self, tmp_path: Path):
        path = tmp_path / "locked.xlsx"
        path.write_bytes(_OLE_HEADER + b"\x00" * 512)

        document = extract_excel(path)

        assert document.is_encrypted is True
        assert document.needs_quarantine is True
        assert document.error == "encrypted_document"

    def test_legacy_binary_xls_is_reported_separately_from_encryption(self, tmp_path: Path):
        path = tmp_path / "old.xls"
        path.write_bytes(_OLE_HEADER + b"\x00" * 512)

        document = extract_excel(path)

        assert document.is_encrypted is False
        assert document.is_valid is False
        assert document.error == "unsupported_legacy_xls"

    def test_zip_bomb_is_rejected_before_the_parser_runs(self, tmp_path: Path):
        path = tmp_path / "bomb.xlsx"
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            archive.writestr("xl/worksheets/sheet1.xml", "a" * 1_000_000)

        document = extract_excel(path)

        assert document.is_valid is False
        assert document.error == "zip_compression_ratio_limit_exceeded"

    def test_cell_budget_stops_reading_and_flags_truncation(self, tmp_path: Path, monkeypatch):
        path = _write_workbook(
            tmp_path / "big.xlsx",
            [("Sheet1", [[f"행{index}", index] for index in range(10)])],
        )
        monkeypatch.setattr(excel_extractor, "_MAX_CELLS_TOTAL", 6)

        document = extract_excel(path)

        assert document.truncated is True
        assert document.is_valid is True
        assert len(document.sheets[0].rows) == 3
