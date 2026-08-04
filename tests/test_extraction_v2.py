from __future__ import annotations

import hashlib
from pathlib import Path

import pymupdf
import pytest
from openpyxl import Workbook

from rd2.extraction import excel_text, hwp_text, storage
from rd2.extraction.excel_text import extract_excel_document
from rd2.extraction.hwp_text import extract_hwp_document
from rd2.extraction.pdf_text import _physical_line_payload, extract_pdf_document
from rd2.extraction.pipeline import (
    OCR_QUEUE_MANIFEST_NAME,
    RUN_MANIFEST_NAME,
    extraction_metadata_for,
    iter_source_documents,
    process_document,
    run_extraction,
)
from rd2.extraction.storage import (
    SCHEMA_VERSION,
    build_extraction_id,
    compute_source_sha256,
    extraction_output_path,
    is_current_extraction,
    read_json_gz,
    write_json_gz_atomic,
)
from rd2.extractors.excel import ExtractedExcelDocument
from rd2.extractors.hwp import ExtractedHwpDocument

import extract_documents as extraction_cli  # noqa: E402


def _source_path(tmp_path: Path, name: str) -> tuple[Path, Path]:
    data_root = tmp_path / "data"
    path = data_root / "moe" / "report" / name
    path.parent.mkdir(parents=True, exist_ok=True)
    return data_root, path


def _make_two_line_pdf(path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=200, height=160)
    page.insert_text((20, 40), "first physical line", fontsize=10)
    page.insert_text((20, 58), "second physical line", fontsize=10)
    document.save(path)
    document.close()


def _make_blank_pdf(path: Path) -> None:
    document = pymupdf.open()
    document.new_page(width=200, height=160)
    document.save(path)
    document.close()


def _make_pdf_with_page_texts(path: Path, page_texts: list[str | None]) -> None:
    document = pymupdf.open()
    for text in page_texts:
        page = document.new_page(width=200, height=160)
        if text is not None:
            page.insert_text((20, 40), text, fontsize=10)
    document.save(path)
    document.close()


def _make_encrypted_pdf(path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=200, height=160)
    page.insert_text((20, 40), "encrypted text")
    document.save(
        path,
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner-secret",
        user_pw="user-secret",
    )
    document.close()


def test_pdf_v2_schema_stores_physical_lines_and_rounded_geometry(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "42_sample.pdf")
    _make_two_line_pdf(source_path)

    payload = extract_pdf_document(source_path, data_root=data_root)

    assert set(payload) == {
        "schema_version",
        "extraction_id",
        "source_sha256",
        "source_path",
        "source",
        "doc_type",
        "doc_id",
        "source_format",
        "extraction",
        "status",
        "error",
        "quality",
        "pages",
    }
    assert payload["schema_version"] == 2
    assert payload["source_path"] == "data/moe/report/42_sample.pdf"
    assert payload["source"] == "moe"
    assert payload["doc_type"] == "report"
    assert payload["doc_id"] == "42"
    assert payload["source_format"] == "pdf"
    assert payload["status"] == "ok"
    assert payload["error"] is None
    assert set(payload["quality"]) == {
        "has_text_layer",
        "needs_ocr",
        "needs_quarantine",
        "pages_needing_ocr",
        "avg_chars_per_page",
        "sparse_page_count",
        "sparse_page_ratio",
        "ocr_severity",
        "warnings",
    }
    assert payload["quality"]["sparse_page_count"] == 0
    assert payload["quality"]["sparse_page_ratio"] == 0.0
    assert payload["quality"]["ocr_severity"] == "clean"

    page = payload["pages"][0]
    assert set(page) == {"page", "width_pt", "height_pt", "rotation", "lines"}
    assert page["page"] == 1
    assert page["width_pt"] == 200.0
    assert page["height_pt"] == 160.0
    assert [line["text"] for line in page["lines"]] == [
        "first physical line",
        "second physical line",
    ]
    assert [line["line_id"] for line in page["lines"]] == [0, 1]
    assert [line["order"] for line in page["lines"]] == [0, 1]
    for line in page["lines"]:
        assert set(line) == {
            "line_id",
            "block_id",
            "order",
            "text",
            "bbox_pt",
            "style_runs",
        }
        assert all(value == round(value, 1) for value in line["bbox_pt"])
        assert line["style_runs"][0]["start"] == 0
        assert line["style_runs"][-1]["end"] == len(line["text"])


def test_pdf_style_runs_coalesce_adjacent_identical_styles_and_decode_flags():
    raw_line = {
        "spans": [
            {
                "text": "AB",
                "bbox": [1.04, 2.06, 3.04, 4.06],
                "font": "Example",
                "size": 10.04,
                "flags": 16,
                "color": 0,
            },
            {
                "text": "CD",
                "bbox": [3.04, 2.06, 5.04, 4.06],
                "font": "Example",
                "size": 10.02,
                "flags": 16,
                "color": 0,
            },
            {
                "text": "E",
                "bbox": [5.04, 2.06, 6.04, 4.06],
                "font": "Example",
                "size": 10.16,
                "flags": 2,
                "color": 7,
            },
        ]
    }

    line = _physical_line_payload(raw_line, line_id=7, block_id=3, order=2)

    assert line is not None
    assert line["text"] == "ABCDE"
    assert line["bbox_pt"] == [1.0, 2.1, 6.0, 4.1]
    assert line["style_runs"] == [
        {
            "start": 0,
            "end": 4,
            "font": "Example",
            "size_pt": 10.0,
            "bold": True,
            "italic": False,
            "color": 0,
        },
        {
            "start": 4,
            "end": 5,
            "font": "Example",
            "size_pt": 10.2,
            "bold": False,
            "italic": True,
            "color": 7,
        },
    ]


def test_blank_pdf_is_a_successful_needs_ocr_snapshot(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "blank.pdf")
    _make_blank_pdf(source_path)

    payload = extract_pdf_document(source_path, data_root=data_root)

    assert payload["status"] == "needs_ocr"
    assert payload["error"] is None
    assert payload["quality"]["has_text_layer"] is False
    assert payload["quality"]["needs_ocr"] is True
    assert payload["quality"]["needs_quarantine"] is False
    assert payload["quality"]["pages_needing_ocr"] == [1]
    assert payload["quality"]["sparse_page_count"] == 1
    assert payload["quality"]["sparse_page_ratio"] == 1.0
    assert payload["quality"]["ocr_severity"] == "image_pdf_candidate"
    assert payload["quality"]["warnings"] == ["image_pdf_candidate"]
    assert payload["pages"][0]["lines"] == []


def test_pdf_with_less_than_half_sparse_pages_stays_readable(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "partial.pdf")
    _make_pdf_with_page_texts(
        source_path,
        [None, "enough text on page two", "enough text on page three"],
    )

    payload = extract_pdf_document(source_path, data_root=data_root)

    assert payload["status"] == "ok"
    assert payload["quality"]["needs_ocr"] is False
    assert payload["quality"]["pages_needing_ocr"] == [1]
    assert payload["quality"]["sparse_page_ratio"] == 0.3333
    assert payload["quality"]["ocr_severity"] == "partial_text"
    assert payload["quality"]["warnings"] == ["partial_text_layer"]


def test_pdf_with_half_sparse_pages_needs_ocr(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "majority-threshold.pdf")
    _make_pdf_with_page_texts(
        source_path,
        [None, None, "enough text on page three", "enough text on page four"],
    )

    payload = extract_pdf_document(source_path, data_root=data_root)

    assert payload["status"] == "needs_ocr"
    assert payload["quality"]["needs_ocr"] is True
    assert payload["quality"]["sparse_page_ratio"] == 0.5
    assert payload["quality"]["ocr_severity"] == "needs_ocr"
    assert payload["quality"]["warnings"] == ["majority_sparse_pages"]


def test_pdf_with_ninety_percent_sparse_pages_is_image_candidate(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "image-candidate.pdf")
    _make_pdf_with_page_texts(
        source_path,
        [None] * 9 + ["one page still has enough extractable text"],
    )

    payload = extract_pdf_document(source_path, data_root=data_root)

    assert payload["status"] == "needs_ocr"
    assert payload["quality"]["has_text_layer"] is True
    assert payload["quality"]["sparse_page_ratio"] == 0.9
    assert payload["quality"]["ocr_severity"] == "image_pdf_candidate"
    assert payload["quality"]["warnings"] == ["image_pdf_candidate"]


def test_encrypted_pdf_is_serialized_as_quarantined_error(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "encrypted.pdf")
    _make_encrypted_pdf(source_path)

    payload = extract_pdf_document(source_path, data_root=data_root)

    assert payload["status"] == "error"
    assert "encrypted" in payload["error"]
    assert payload["quality"]["needs_quarantine"] is True
    assert payload["quality"]["warnings"] == ["encrypted"]
    assert payload["pages"] == []


def test_hwp_v2_uses_one_logical_page_with_null_geometry(tmp_path: Path, monkeypatch):
    data_root, source_path = _source_path(tmp_path, "77_sample.hwpx")
    source_path.write_bytes(b"placeholder")
    calls: list[bool] = []

    def fake_extract(path: Path, *, include_tables: bool = True):
        calls.append(include_tables)
        return ExtractedHwpDocument(source_path=path, text="첫 문단\n\n둘째 문단")

    monkeypatch.setattr(hwp_text, "extract_hwp", fake_extract)

    payload = extract_hwp_document(source_path, data_root=data_root)

    assert calls == [False]
    assert payload["schema_version"] == 2
    assert payload["source_format"] == "hwpx"
    assert payload["status"] == "ok"
    assert set(payload["quality"]) == {
        "has_text_layer",
        "needs_ocr",
        "needs_quarantine",
        "pages_needing_ocr",
        "avg_chars_per_page",
        "sparse_page_count",
        "sparse_page_ratio",
        "ocr_severity",
        "warnings",
    }
    assert payload["quality"]["sparse_page_count"] == 0
    assert payload["quality"]["sparse_page_ratio"] == 0.0
    assert payload["quality"]["ocr_severity"] == "clean"
    page = payload["pages"][0]
    assert (page["width_pt"], page["height_pt"], page["rotation"]) == (None, None, None)
    assert [line["text"] for line in page["lines"]] == ["첫 문단", "둘째 문단"]
    assert [line["line_id"] for line in page["lines"]] == [0, 1]
    assert all(line["block_id"] is None for line in page["lines"])
    assert all(line["bbox_pt"] is None for line in page["lines"])
    assert all(line["style_runs"] == [] for line in page["lines"])


def test_hwp_v2_marks_little_or_no_text_as_needing_ocr(tmp_path: Path, monkeypatch):
    data_root, source_path = _source_path(tmp_path, "empty.hwp")
    source_path.write_bytes(b"placeholder")
    monkeypatch.setattr(
        hwp_text,
        "extract_hwp",
        lambda path, **_kwargs: ExtractedHwpDocument(source_path=path, text=" "),
    )

    payload = extract_hwp_document(source_path, data_root=data_root)

    assert payload["status"] == "needs_ocr"
    assert payload["error"] is None
    assert payload["quality"]["has_text_layer"] is False
    assert payload["quality"]["needs_ocr"] is True
    assert payload["quality"]["needs_quarantine"] is False
    assert payload["quality"]["pages_needing_ocr"] == [1]
    assert payload["quality"]["sparse_page_count"] == 1
    assert payload["quality"]["sparse_page_ratio"] == 1.0
    assert payload["quality"]["ocr_severity"] == "needs_ocr"
    assert payload["quality"]["warnings"] == ["little_or_no_text"]
    assert payload["pages"][0]["lines"] == []


def test_hwp_v2_splits_oversized_logical_lines_without_losing_text(
    tmp_path: Path,
    monkeypatch,
):
    data_root, source_path = _source_path(tmp_path, "long-table-row.hwpx")
    source_path.write_bytes(b"placeholder")
    long_text = "가" * (hwp_text._MAX_LOGICAL_LINE_CHARS + 3)
    monkeypatch.setattr(
        hwp_text,
        "extract_hwp",
        lambda path, **_kwargs: ExtractedHwpDocument(source_path=path, text=long_text),
    )

    payload = extract_hwp_document(source_path, data_root=data_root)
    lines = payload["pages"][0]["lines"]

    assert payload["status"] == "ok"
    assert [len(line["text"]) for line in lines] == [
        hwp_text._MAX_LOGICAL_LINE_CHARS,
        3,
    ]
    assert "".join(line["text"] for line in lines) == long_text
    assert payload["quality"]["warnings"] == ["oversized_logical_line_split"]


def test_encrypted_hwp_is_quarantined_without_fake_page(tmp_path: Path, monkeypatch):
    data_root, source_path = _source_path(tmp_path, "encrypted.hwp")
    source_path.write_bytes(b"placeholder")
    monkeypatch.setattr(
        hwp_text,
        "extract_hwp",
        lambda path, **_kwargs: ExtractedHwpDocument(
            source_path=path,
            is_encrypted=True,
            error="encrypted",
        ),
    )

    payload = extract_hwp_document(source_path, data_root=data_root)

    assert payload["status"] == "quarantine"
    assert payload["error"] == "encrypted"
    assert payload["quality"]["needs_quarantine"] is True
    assert payload["pages"] == []


def _make_workbook(path: Path, sheets: list[tuple[str, list[list]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        worksheet = workbook.create_sheet(title=name)
        for row in rows:
            worksheet.append(row)
    workbook.save(path)
    workbook.close()


def test_excel_v2_maps_each_worksheet_to_a_page_of_markdown_rows(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "31_활동내역.xlsx")
    _make_workbook(
        source_path,
        [
            ("2024년 1차", [["이사", "회차"], ["홍길동", 3]]),
            ("2024년 2차", [["이사"], ["김철수"]]),
        ],
    )

    payload = extract_excel_document(source_path, data_root=data_root)

    assert payload["schema_version"] == 2
    assert payload["source_format"] == "xlsx"
    assert payload["doc_id"] == "31"
    assert payload["status"] == "ok"
    assert payload["error"] is None
    assert payload["extraction"]["extractor"] == "openpyxl"
    assert [page["page"] for page in payload["pages"]] == [1, 2]

    first_page = payload["pages"][0]
    assert set(first_page) == {"page", "width_pt", "height_pt", "rotation", "lines"}
    assert (first_page["width_pt"], first_page["height_pt"], first_page["rotation"]) == (
        None,
        None,
        None,
    )
    assert [line["text"] for line in first_page["lines"]] == [
        "2024년 1차",
        "| 이사 | 회차 |",
        "| 홍길동 | 3 |",
    ]
    # 시트 이름 줄은 어느 행에도 속하지 않고, 행에서 나온 줄만 행 번호로 묶인다.
    assert [line["block_id"] for line in first_page["lines"]] == [None, 0, 1]
    assert all(line["bbox_pt"] is None for line in first_page["lines"])
    assert all(line["style_runs"] == [] for line in first_page["lines"])
    # line_id는 워크북 전체에서 이어지고, order는 페이지 안에서 다시 0부터 센다.
    assert [line["line_id"] for line in payload["pages"][1]["lines"]] == [3, 4, 5]
    assert [line["order"] for line in payload["pages"][1]["lines"]] == [0, 1, 2]


def test_excel_v2_wraps_an_oversized_cell_without_losing_text(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "memo.xlsx")
    long_text = "가" * (excel_text._MAX_CELL_CHARS + 5)
    _make_workbook(source_path, [("Sheet1", [[long_text]])])

    payload = extract_excel_document(source_path, data_root=data_root)
    lines = payload["pages"][0]["lines"][1:]  # 시트 이름 줄 제외

    assert payload["status"] == "ok"
    assert payload["quality"]["warnings"] == ["oversized_cell_split"]
    assert all(len(line["text"]) <= excel_text._MAX_LOGICAL_LINE_CHARS for line in lines)
    assert "".join(line["text"].strip("| ") for line in lines) == long_text
    assert all(line["block_id"] == 0 for line in lines)


def test_empty_workbook_is_a_successful_needs_ocr_snapshot(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "empty.xlsx")
    _make_workbook(source_path, [("Sheet1", [])])

    payload = extract_excel_document(source_path, data_root=data_root)

    # 시트 이름 줄만 남은 워크북은 "내용 있음"이 아니다 — 시트는 비어 있어도
    # 항상 이름을 갖기 때문이다.
    assert payload["status"] == "needs_ocr"
    assert payload["error"] is None
    assert payload["quality"]["has_text_layer"] is False
    assert payload["quality"]["needs_quarantine"] is False
    assert payload["quality"]["pages_needing_ocr"] == [1]
    assert payload["quality"]["warnings"] == ["little_or_no_text"]
    assert [line["text"] for line in payload["pages"][0]["lines"]] == ["Sheet1"]


def test_encrypted_workbook_is_quarantined_without_fake_page(tmp_path: Path, monkeypatch):
    data_root, source_path = _source_path(tmp_path, "locked.xlsx")
    source_path.write_bytes(b"placeholder")
    monkeypatch.setattr(
        excel_text,
        "extract_excel",
        lambda path: ExtractedExcelDocument(
            source_path=path,
            is_encrypted=True,
            error="encrypted_document",
        ),
    )

    payload = extract_excel_document(source_path, data_root=data_root)

    assert payload["status"] == "quarantine"
    assert payload["error"] == "encrypted_document"
    assert payload["quality"]["needs_quarantine"] is True
    assert payload["quality"]["warnings"] == ["encrypted"]
    assert payload["pages"] == []


def test_pipeline_routes_workbooks_to_the_openpyxl_extractor(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "42_현황.xlsx")
    extracted_root = data_root / "extracted"
    _make_workbook(source_path, [("Sheet1", [["항목", "값"]])])

    assert extraction_metadata_for(source_path)["extractor"] == "openpyxl"
    assert [path.name for path in iter_source_documents(data_root, source_format="xlsx")] == [
        "42_현황.xlsx"
    ]

    action, output_path, payload = process_document(
        source_path,
        data_root=data_root,
        extracted_root=extracted_root,
    )

    assert action == "processed"
    assert output_path.name == "42_현황.xlsx.json.gz"
    assert payload["status"] == "ok"
    assert read_json_gz(output_path)["pages"][0]["lines"][-1]["text"] == "| 항목 | 값 |"


def test_output_path_preserves_source_extension_for_same_stem(tmp_path: Path):
    data_root = tmp_path / "data"
    extracted_root = data_root / "extracted"
    pdf = data_root / "moe" / "report" / "same.pdf"
    hwp = data_root / "moe" / "report" / "same.hwp"

    assert extraction_output_path(pdf, data_root, extracted_root).name == "same.pdf.json.gz"
    assert extraction_output_path(hwp, data_root, extracted_root).name == "same.hwp.json.gz"
    assert extraction_output_path(pdf, data_root, extracted_root) != extraction_output_path(
        hwp, data_root, extracted_root
    )


def test_gzip_json_roundtrip_is_compact_and_deterministic(tmp_path: Path):
    payload = {"schema_version": 2, "text": "한글", "nested": {"ok": True}}
    first = tmp_path / "first.json.gz"
    second = tmp_path / "second.json.gz"

    write_json_gz_atomic(first, payload)
    write_json_gz_atomic(second, payload)

    assert read_json_gz(first) == payload
    assert first.read_bytes() == second.read_bytes()


def test_atomic_write_failure_preserves_existing_target(tmp_path: Path, monkeypatch):
    target = tmp_path / "document.json.gz"
    original = {"version": "old"}
    write_json_gz_atomic(target, original)

    def fail_replace(source, destination):
        raise OSError("simulated replace failure")

    monkeypatch.setattr(storage.os, "replace", fail_replace)
    with pytest.raises(OSError, match="simulated replace failure"):
        write_json_gz_atomic(target, {"version": "new"})

    assert read_json_gz(target) == original
    assert list(tmp_path.glob(f".{target.name}.*.tmp")) == []


def test_extraction_identity_and_current_cache_invalidate_on_source_or_config_change(tmp_path: Path):
    source = tmp_path / "source.pdf"
    source.write_bytes(b"source bytes")
    source_digest = compute_source_sha256(source)
    assert source_digest == hashlib.sha256(b"source bytes").hexdigest()

    extraction = {
        "profile": "layout-lite-v1",
        "extractor": "test",
        "extractor_version": "1",
        "config": {"precision": 1},
    }
    extraction_id = build_extraction_id(source_digest, extraction)
    changed_id = build_extraction_id(source_digest, {**extraction, "config": {"precision": 2}})
    assert extraction_id != changed_id

    output = tmp_path / "source.pdf.json.gz"
    write_json_gz_atomic(
        output,
        {
            "schema_version": SCHEMA_VERSION,
            "source_sha256": source_digest,
            "extraction_id": extraction_id,
        },
    )
    assert is_current_extraction(output, source_digest)
    assert is_current_extraction(output, source_digest, expected_extraction_id=extraction_id)
    assert not is_current_extraction(output, source_digest, expected_extraction_id=changed_id)
    assert not is_current_extraction(output, "0" * 64)


def test_process_document_reextracts_when_source_content_changes(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "broken.pdf")
    extracted_root = data_root / "extracted"
    source_path.write_bytes(b"not a pdf")

    first_action, output_path, first = process_document(
        source_path,
        data_root=data_root,
        extracted_root=extracted_root,
    )
    second_action, _, second = process_document(
        source_path,
        data_root=data_root,
        extracted_root=extracted_root,
    )
    source_path.write_bytes(b"changed, still not a pdf")
    third_action, _, third = process_document(
        source_path,
        data_root=data_root,
        extracted_root=extracted_root,
    )

    assert first_action == "processed"
    assert second_action == "skipped"
    assert third_action == "processed"
    assert first["source_sha256"] == second["source_sha256"]
    assert third["source_sha256"] != first["source_sha256"]
    assert read_json_gz(output_path)["source_sha256"] == third["source_sha256"]


def test_run_manifest_is_atomic_and_records_partial_failures(tmp_path: Path):
    data_root, source_path = _source_path(tmp_path, "invalid.pdf")
    extracted_root = data_root / "extracted"
    source_path.write_bytes(b"not a pdf")

    manifest = run_extraction(
        [source_path],
        data_root=data_root,
        extracted_root=extracted_root,
    )

    assert manifest["status"] == "partial"
    assert manifest["counts"] == {
        "total": 1,
        "processed": 1,
        "skipped": 0,
        "succeeded": 0,
        "failed": 1,
    }
    assert manifest["failures"][0]["source_path"] == "data/moe/report/invalid.pdf"
    assert manifest["artifacts"] == []
    assert read_json_gz(extracted_root / RUN_MANIFEST_NAME) == manifest


def test_run_manifest_records_deterministic_success_and_skip_artifacts(tmp_path: Path):
    data_root, first_source = _source_path(tmp_path, "a.pdf")
    _, second_source = _source_path(tmp_path, "b.pdf")
    extracted_root = data_root / "extracted"
    _make_two_line_pdf(first_source)
    _make_two_line_pdf(second_source)

    first_manifest = run_extraction(
        [second_source, first_source],
        data_root=data_root,
        extracted_root=extracted_root,
    )

    assert first_manifest["status"] == "complete"
    assert first_manifest["counts"]["succeeded"] == 2
    assert [entry["source_path"] for entry in first_manifest["artifacts"]] == [
        "data/moe/report/a.pdf",
        "data/moe/report/b.pdf",
    ]
    assert [entry["output_path"] for entry in first_manifest["artifacts"]] == [
        "moe/report/a.pdf.json.gz",
        "moe/report/b.pdf.json.gz",
    ]
    assert all(entry["status"] == "ok" for entry in first_manifest["artifacts"])
    assert all(entry["extraction_id"] for entry in first_manifest["artifacts"])

    second_manifest = run_extraction(
        [second_source, first_source],
        data_root=data_root,
        extracted_root=extracted_root,
    )

    assert second_manifest["counts"]["skipped"] == 2
    assert second_manifest["artifacts"] == first_manifest["artifacts"]


def test_run_extraction_writes_reference_only_ocr_queue_manifest(tmp_path: Path):
    data_root, blank_source = _source_path(tmp_path, "blank.pdf")
    _, readable_source = _source_path(tmp_path, "readable.pdf")
    extracted_root = data_root / "extracted"
    ocr_queue_root = data_root / "ocr_queue"
    _make_blank_pdf(blank_source)
    _make_two_line_pdf(readable_source)

    manifest = run_extraction(
        [readable_source, blank_source],
        data_root=data_root,
        extracted_root=extracted_root,
        ocr_queue_root=ocr_queue_root,
    )

    queue = read_json_gz(ocr_queue_root / OCR_QUEUE_MANIFEST_NAME)
    assert queue["source_run_id"] == manifest["run_id"]
    assert queue["source_run_status"] == "complete"
    assert queue["count"] == 1
    assert queue["entries"] == [
        {
            "source_path": "data/moe/report/blank.pdf",
            "output_path": "moe/report/blank.pdf.json.gz",
            "extraction_id": manifest["artifacts"][0]["extraction_id"],
            "source_format": "pdf",
            "has_text_layer": False,
            "sparse_page_ratio": 1.0,
            "ocr_severity": "image_pdf_candidate",
            "pages_needing_ocr": [1],
        }
    ]
    assert list(ocr_queue_root.glob("*.pdf.json.gz")) == []


def test_cli_limit_slices_deterministically_ordered_documents(tmp_path: Path, monkeypatch):
    data_root = tmp_path / "data"
    extracted_root = data_root / "extracted"
    source_root = data_root / "moe" / "report"
    source_root.mkdir(parents=True)
    for name in ("c.pdf", "a.pdf", "b.pdf"):
        (source_root / name).write_bytes(b"placeholder")

    captured: list[Path] = []

    def fake_run(documents, **kwargs):
        captured.extend(documents)
        return {
            "status": "complete",
            "counts": {"total": 2, "processed": 2, "skipped": 0, "succeeded": 2, "failed": 0},
            "failures": [],
        }

    monkeypatch.setattr(extraction_cli, "_DATA_ROOT", data_root)
    monkeypatch.setattr(extraction_cli, "_EXTRACTED_ROOT", extracted_root)
    monkeypatch.setattr(extraction_cli, "run_extraction", fake_run)

    assert extraction_cli.main(["--source", "moe", "--format", "pdf", "--limit", "2"]) == 0
    assert [path.name for path in captured] == ["a.pdf", "b.pdf"]


def test_cli_refuses_empty_scope_before_replacing_manifest(tmp_path: Path, monkeypatch):
    data_root = tmp_path / "data"
    data_root.mkdir()
    called = False

    def fake_run(*_args, **_kwargs):
        nonlocal called
        called = True
        raise AssertionError("empty scope must not publish a manifest")

    monkeypatch.setattr(extraction_cli, "_DATA_ROOT", data_root)
    monkeypatch.setattr(extraction_cli, "_EXTRACTED_ROOT", data_root / "extracted")
    monkeypatch.setattr(extraction_cli, "run_extraction", fake_run)

    with pytest.raises(SystemExit):
        extraction_cli.main(["--source", "missing"])
    assert called is False


def test_source_scope_cannot_escape_data_root(tmp_path: Path):
    data_root = tmp_path / "data"
    data_root.mkdir()

    with pytest.raises(ValueError, match="escapes data root"):
        list(iter_source_documents(data_root, source="../outside"))


def test_cli_partial_exit_requires_allow_partial(tmp_path: Path, monkeypatch):
    data_root, source_path = _source_path(tmp_path, "invalid.pdf")
    source_path.write_bytes(b"placeholder")
    extracted_root = data_root / "extracted"

    def fake_run(documents, **kwargs):
        return {
            "status": "partial",
            "counts": {"total": 1, "processed": 1, "skipped": 0, "succeeded": 0, "failed": 1},
            "failures": [{"source_path": str(source_path), "error": "broken"}],
        }

    monkeypatch.setattr(extraction_cli, "_DATA_ROOT", data_root)
    monkeypatch.setattr(extraction_cli, "_EXTRACTED_ROOT", extracted_root)
    monkeypatch.setattr(extraction_cli, "run_extraction", fake_run)

    assert extraction_cli.main(["--document", str(source_path)]) == 1
    assert extraction_cli.main(["--document", str(source_path), "--allow-partial"]) == 0
